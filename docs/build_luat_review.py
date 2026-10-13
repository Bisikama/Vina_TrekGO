from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).parent / 'LUAT_GPS_MAP_NOTIFICATION_MVP_2026-10-12.docx'
d = Document()
s = d.sections[0]
s.page_height, s.page_width = Inches(11.69), Inches(8.27)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
 st = d.styles[name]
 st.font.name = 'Arial'
 st.font.color.rgb = RGBColor(0,0,0)
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.space_after = Pt(7)
d.styles['Normal'].paragraph_format.line_spacing = 1.12
d.styles['Title'].font.size = Pt(23)
d.styles['Heading 1'].font.size = Pt(16)
d.styles['Heading 2'].font.size = Pt(12)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(head, rows, widths):
 t=d.add_table(rows=1, cols=len(head)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,head): c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row): c.text=v
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr()
  keep=OxmlElement('w:cantSplit');trpr.append(keep)
  if i==0: trpr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Inches(w)
   pr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if i==0 else 'FFFFFF');pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2)
    for r in para.runs:r.font.size=Pt(9);r.bold=i==0
 p('')

d.add_heading('GPS bản đồ và thông báo trong MVP của Luật',0)
p('TrekGo • Phân tích triển khai và danh sách công việc • 12 tháng 10 năm 2026')
p('Mục tiêu bàn giao cho Luật là một luồng trekking chạy xuyên suốt: mở chuyến đi, thấy tuyến và checkpoint thật, lấy vị trí điện thoại, ghi đường đã đi, đồng bộ GPS, cảnh báo lệch tuyến và ghi nhận đến checkpoint. Hiện nền tảng bản đồ web và API đọc ngữ cảnh đã có, còn luồng GPS mobile và hệ thống thông báo cần được nối hoàn chỉnh.')
h('Kết luận về hiện trạng')
table(['Hạng mục','Kết quả đối chiếu mã nguồn'],[
 ('Map web','Có MapLibre GL JS, ảnh vệ tinh Esri, terrain DEM Terrarium, polyline và marker checkpoint.'),
 ('GPS mobile','Màn TripMap vẫn là BẢN ĐỒ MÔ PHỎNG; dữ liệu tiến độ và cảnh báo đến từ mockData. Chưa có thư viện lấy vị trí trong package.json.'),
 ('Backend GPS','Có GET navigation-context và GET map-state; schema có session, location, deviation, checkpoint event và outbox. Chưa có các API ghi GPS trong module geo.'),
 ('Notification','Có schema PostgreSQL và định hướng outbox; chưa có module consumer, API hộp thư hay tích hợp push trong nguồn hiện tại.'),
 ('Weather','Có WeatherSnapshot trong schema; rental đã đọc snapshot. Chưa có module lấy dữ liệu thời tiết từ provider.')
],[1.35,5.35])
h('Căn cứ giao việc')
p('api/docs/luat.md ghi rõ Luật sở hữu Area C gồm auth integration, map, GPS, deviation, checkpoint execution và weather. Lịch sử Git gắn commit 11efd51 với phần bản đồ 3D và 51cf2dd với geo/navigation cho tác giả luatluat304@gmail.com. Đây là căn cứ nhận diện phần việc, không đồng nghĩa mọi dòng mã hiện tại đều do riêng Luật viết.')
p('Bản đối chiếu lấy working tree tại HEAD 811267a làm mốc. Trạng thái “có mã” chưa phải nghiệm thu chạy thực tế. Danh sách công việc bên dưới là đề xuất triển khai; các API được đánh dấu đề xuất không phải endpoint đã tồn tại.')

page('1 Cách bản đồ hiện tại được xây dựng')
h('Ba lớp đang phối hợp')
p('Lớp dữ liệu lấy geometry và checkpoint từ backend. Ví dụ TripRouteReviewMap gọi /api/v1/system-trails/:trailCode, nhận routeGeojson, checkpoints và thông tin camera. Đây là tọa độ tuyến lưu sẵn, chưa phải vị trí GPS trực tiếp của người đang cầm điện thoại. [S2, S3]')
p('Lớp nền dùng ảnh vệ tinh Esri World Imagery. Lớp độ cao dùng raster-dem mã hóa Terrarium; khu vực Bà Đen có thể dùng /tiles/dem/{z}/{x}/{y}.png, khu vực khác dùng elevation-tiles-prod trên S3. BaseMap3D gọi setTerrain, thêm hillshade, cho phép xoay/nghiêng và bật tắt 3D. Worker hiện được trỏ đến unpkg bản 6.11.2. [S2]')
p('Lớp hiển thị nghiệp vụ dùng drawTrailPolyline: chuẩn hóa LineString/MultiLineString thành GeoJSON, tạo source và ba lớp đường gồm bóng, viền, nét chính; source đã có được cập nhật bằng setData. renderCheckpointMarkers dùng Marker và Popup tại [longitude, latitude]. Các helper chờ style sẵn sàng và dọn layer/source hỗ trợ tái sử dụng. [S2]')
h('Cách sử dụng lại cho màn theo dõi chuyến đi')
p('1. Gọi navigation-context của trip đang tham gia, lấy activeTrail.routeGeojson và checkpoint. Không dùng mặc định một trailCode làm nguồn chân lý cho hành trình đang diễn ra.')
p('2. Trên web, dựng BaseMap3D, chờ onMapReady hoặc runWhenMapStyleReady rồi vẽ tuyến kế hoạch và checkpoint. Tuyến thực tế phải có source/id riêng để không ghi đè tuyến kế hoạch.')
p('3. Nguồn vị trí thiết bị cập nhật current marker và polyline thực tế. Khi có điểm mới, cập nhật dữ liệu source; không dựng lại toàn bộ map mỗi lần nhận GPS. Quy ước GeoJSON là [lng, lat], trong khi payload nghiệp vụ có thể là latitude/longitude.')
p('4. Khi đổi tuyến, đọc lại navigation-context và thay active route; giữ đường thực tế và checkpoint của trip. Khi rời màn hình, dọn map/listener; việc dừng thu GPS phải theo trạng thái session, không phụ thuộc riêng việc đóng màn bản đồ.')
h('Giới hạn phải hiểu đúng')
p('MapLibre GL JS trong repo là thư viện web dùng DOM/WebGL. GPS cần nguồn thu vị trí riêng. BaseMap3D chưa có GeolocateControl hay watchPosition. Độ phóng đại terrain mặc định 2.0 là hiển thị thị giác, không nên dùng làm căn cứ tính độ cao hoặc an toàn hành trình. Các màn map cũ còn tự khởi tạo MapLibre, nên tái sử dụng helper chung khi nối tracking. [S2, S3]')

page('2 Luồng GPS đề xuất cho mobile')
p('Mobile dùng Expo ~57.0.23 và React Native 0.86.3. Đề xuất dùng expo-location để thu vị trí, chọn adapter bản đồ native sau một bản thử nhỏ trên thiết bị thật. Đây là phần bổ sung, chưa có trong mobile/package.json. [S4]')
h('Thu vị trí và vẽ bản đồ')
p('Với Expo SDK 57, xin quyền bằng requestForegroundPermissionsAsync và theo dõi bằng watchPositionAsync. API theo dõi này chỉ cập nhật khi app ở foreground; muốn chạy nền phải thiết kế riêng startLocationUpdatesAsync, quyền background và vòng đời task. Cài bằng npx expo install expo-location để chọn bản phù hợp SDK. [W1]')
p('Ứng viên hiển thị native là @maplibre/maplibre-react-native. Thư viện bọc MapLibre Native, cần setup và rebuild app; không thể chép nguyên BaseMap3D.tsx vốn dùng thẻ div sang React Native. Cần thử route, marker và terrain trên Android/iOS mục tiêu trước khi chốt khả năng 3D; việc nghiêng camera không đủ chứng minh terrain 3D tương đương web. [W2]')
h('Chuỗi xử lý từ thiết bị đến backend')
p('Quyền vị trí → GPS session → mẫu vị trí hợp lệ → hàng đợi local → marker và actual path → batch API → lưu DB → kiểm tra lệch tuyến/checkpoint → cập nhật trạng thái và phát sự kiện.')
p('Mỗi mẫu nên có clientSampleId ổn định, recordedAt, latitude, longitude, accuracyM và các trường speedMps/headingDeg/altitudeM nếu có. Backend ánh xạ user từ token, kiểm tra chủ session, trạng thái trip và ngưỡng accuracy theo GpsPolicy. Không tin userId tự gửi từ client. [S5]')
p('Thu mẫu và gửi mẫu là hai nhịp riêng. Khoảng đồng bộ 10–30 giây là ví dụ từ đặc tả, chưa phải cấu hình đang chạy. Vẽ local để phản hồi nhanh; hàng đợi bền vững giữ mẫu khi mất mạng, retry giữ nguyên clientSampleId, server chống trùng theo (sessionId, clientSampleId). Sau stop phải flush mẫu còn lại và trả summary nhất quán. [S1, S5]')
h('Lệch tuyến và checkpoint')
p('Tính khoảng cách ngắn nhất từ điểm GPS tới đoạn tuyến active bằng mét; không chỉ đo tới đỉnh gần nhất. Chỉ chuyển sang DEVIATED sau số mẫu liên tiếp theo policy; có RECOVERED, cooldown và thời điểm vị trí gần nhất. Cảnh báo local có thể hiển thị sớm, backend xác thực và ghi sự kiện cho leader. [S1]')
p('Checkpoint dùng radius từ navigation context. Backend xác minh thời gian, accuracy, vị trí và quyền; ghi một arrival cho user/trip/checkpoint rồi mở mission đúng vai trò. Arrival và mission completion là hai bước khác nhau; Mảng A vẫn sở hữu quy tắc hoàn thành trip. Khi không có vị trí mới, hiển thị “vị trí cũ/chưa xác định”, không tự suy ra an toàn.')

page('3 API và điểm cần hoàn thiện')
p('Tất cả route dưới đây có prefix /api/v1. Hai GET đầu và switch-trail đã có controller; các nhóm còn lại là contract trong tài liệu cần triển khai hoặc thống nhất thêm. [S1, S6]')
table(['Trạng thái','Method và route'],[
 ('Có mã','GET /trips/:tripId/navigation-context'),
 ('Có mã','GET /geo/trips/:tripId/map-state'),
 ('Có mã','POST /trips/:tripId/switch-trail'),
 ('Cần làm','POST /geo/trips/:tripId/sessions\nPOST /geo/sessions/:sessionId/locations/batch'),
 ('Cần làm','GET /geo/sessions/:sessionId/path\nPOST /geo/sessions/:sessionId/stop'),
 ('Cần làm','GET /geo/trips/:tripId/deviations\nGET /geo/trips/:tripId/members/last-locations'),
 ('Cần làm','POST /geo/trips/:tripId/checkpoints/:checkpointId/arrivals\nGET /geo/trips/:tripId/checkpoint-progress'),
 ('Cần chốt owner','POST /geo/checkpoint-events/:eventId/mission-complete'),
 ('Cần làm','GET /weather/trips/:tripId')
],[1.25,5.45])
h('Những vấn đề thấy trực tiếp trong code')
p('Map-state chỉ đọc dữ liệu đã lưu; endpoint này không khởi động GPS hay tính deviation mới. Adapter ghi chú ưu tiên ACTIVE nhưng query thực tế chỉ sort createdAt giảm dần, nên có thể lấy session mới nhất đã dừng. Cần chốt ngữ nghĩa session active và session gần nhất. [S7]')
p('Progress đang đếm checkpointEvent của user, không xét missionStatus và không giới hạn các event theo tập checkpoint enabled dùng làm mẫu số. Do đó completedPercentage hiện phản ánh arrival thô, chưa đủ để kết luận hoàn tất nhiệm vụ hay hoàn tất trip. [S7]')
p('Use case map-state cho phép ADMIN vượt kiểm tra quyền, nhưng adapter trả dữ liệu rỗng sớm nếu admin không phải thành viên/creator/leader. Nếu muốn admin xem dữ liệu thì phải thống nhất scope query; không chỉ sửa guard. [S7]')
p('Switch-trail đang cho OPEN hoặc IN_PROGRESS và trả event ACTIVE_TRAIL_CHANGED trong HTTP response. Chuỗi event trong response không tự tạo realtime subscription cho client. Cần thống nhất điều kiện với Mảng A và cách polling/refresh. [S6]')
p('Geo/navigation vẫn gắn TripErrorFilter cũ. Khi có nhiệm vụ thay đổi error handling, phối hợp chuyển cả client và filter theo shared AppError/ZodValidationPipe/GlobalErrorFilter; không tạo envelope riêng. Tài liệu này không yêu cầu tự refactor các module khác.')

page('4 Notification đang đi theo hướng nào')
h('Thiết kế đã ghi nhận và phần mã đang có')
p('OUTBOX_PATTERN_EXPLAINED.md mô tả business service ghi domain event cùng transaction, consumer định kỳ đọc và tạo thông báo. Ví dụ trong tài liệu dùng MongoDB, metadata.outboxEventId và NOTIFICATION_CONSUMER_ENABLED. Đường dẫn consumer được dẫn trong tài liệu không có trong cây api/src hiện tại. Không nên coi ví dụ đó là tính năng đang hoạt động. [S8]')
p('Schema thực tế có Notification trong PostgreSQL: userId, title, content, notificationType, referenceType/referenceId, status UNREAD và readAt. Chưa có trường sourceEventId/deduplicationKey cho notification. api/package.json không khai báo MongoDB/Mongoose; AppModule chưa đăng ký module notification. [S5, S9]')
p('Booking, payment và rental đã ghi/tiêu thụ một số outbox nghiệp vụ. Ví dụ BookingCancellationService xử lý hoàn tiền rồi đánh dấu BOOKING_CANCELLED là PROCESSED. Điều này không chứng minh thông báo đã gửi. RENTAL_NOTIFICATION_EVENTS_URL mới xuất hiện ở khai báo config; các công tắc notification trên UI là state cục bộ. Chưa có push token, push provider hoặc notification inbox API trong các nguồn được đối chiếu. [S9, S10]')
h('Hướng MVP khuyến nghị')
p('Ưu tiên inbox lưu trong PostgreSQL hiện có, consumer độc lập và UI polling. Luật phát sự kiện GPS và xây màn cảnh báo; team cần chỉ định owner cho notification dùng chung của Booking/Payment/Rental. Không mặc định giao lại nghiệp vụ các domain đó cho Luật.')
p('Luồng đề xuất: nghiệp vụ + outbox commit → consumer xác định recipient → ghi notification và dấu đã xử lý của consumer trong cùng transaction → API inbox → chuông/số chưa đọc → bấm mở đúng trip/booking → đánh dấu đã đọc.')
p('Mỗi consumer cần checkpoint/receipt riêng theo (sourceDomain, eventId, consumerName); mỗi thông báo cần khóa unique theo sự kiện, người nhận và loại/kênh. Không tranh quyền cập nhật trạng thái outbox của payment/refund. Chỉ đọc PENDING sẽ có thể bỏ sót event đã được consumer khác đánh dấu PROCESSED; đọc lại toàn bộ mà thiếu cursor/dedup bền vững cũng chưa đủ. Bổ sung schema bằng migration additive. [S10]')
h('Phân biệt ba hình thức thông báo')
p('Cảnh báo trong màn GPS: phản hồi trực tiếp khi lệch tuyến. Inbox: lưu bền để user/leader xem lại. Push ngoài app: một kênh delivery riêng, cần đăng ký token, quyền, retry, xử lý token hết hạn và deep link. Đề xuất inbox/polling cho MVP foreground; nếu nghiệm thu đòi leader nhận cảnh báo khi app đóng thì push phải được nâng vào phạm vi bắt buộc. Không coi polling là thay thế push trong tình huống đó.')
p('API inbox đề xuất: GET /notifications?cursor=...; GET /notifications/unread-count; PATCH /notifications/:id/read. Mọi thao tác phải giới hạn theo current user. Có retry/backoff và trạng thái lỗi để gửi bù; lỗi notification không làm rollback giao dịch booking đã hoàn thành.')

page('5 Danh sách công việc MVP giao cho Luật')
p('Thứ tự dưới đây ưu tiên hoàn thành luồng thật trước các tính năng mở rộng. P0 là bắt buộc theo phạm vi GPS cốt lõi trong đặc tả; các thay đổi ưu tiên được ghi rõ là đề xuất. [S1]')
table(['Mốc','Đầu việc và tiêu chí nghiệm thu'],[
 ('C0 • P0','Nối auth dùng chung vào mobile và chốt contract hiện hành. Token hợp lệ gọi được API; sai quyền bị từ chối; không tạo auth mới. Backend đã có login/refresh qua adapter Supabase, khác mô tả SDK trực tiếp ban đầu.'),
 ('C1 • P0','Thay map mock bằng map thật. Trip được cấp quyền hiển thị active route/checkpoint từ API; đổi tuyến cập nhật đúng; có loading/error/empty state. Thử khả năng 3D trên thiết bị mục tiêu.'),
 ('C2 • P0','Start/batch/path/stop và thu vị trí. Hai lần start đồng thời không tạo hai session active; retry batch không nhân đôi; vị trí hiện tại và actual path dùng mẫu thật; stop không tiếp tục thu ngoài ý muốn.'),
 ('C3 • P0','Deviation và leader alert. Dưới ngưỡng không cảnh báo; vượt đủ N mẫu mới cảnh báo; hồi phục đóng event; leader đúng trip thấy người, vị trí, khoảng cách và thời gian; không spam.'),
 ('C4 • P0','Arrival và mission. Ngoài radius bị từ chối; retry chỉ có một event; User/Leader thấy mission đúng vai trò; completion chuyển về Mảng A và không nhầm arrival với hoàn tất.'),
 ('C5 • P0','Current weather theo đặc tả Luật. Adapter + normalize + cache + updatedAt; provider lỗi trả trạng thái thiếu/cũ, không làm hỏng trip. Forecast/risk là P1.'),
 ('C6 • P0 đề xuất','Nối GPS event với inbox dùng chung; leader nhận deviation một lần và bấm mở đúng trip. Phối hợp owner notification để kiểm chứng gửi bù và chống trùng.'),
 ('C7 • P0 đề xuất','Mất mạng và vòng đời app. Giữ queue qua restart; đồng bộ bù; hiện thời điểm vị trí cũ. Chốt foreground/background, chạy thử trên máy thật và ghi giới hạn demo.')
],[1.05,5.65])
h('Phần để sau luồng P0')
p('P1: thống kê tiến độ nâng cao, history, forecast/risk và bản đồ last-known của cả đoàn. P2/mở rộng: live location chính xác từng giây, chat/mesh offline, offline map đầy đủ, tạo personal trail từ recording. Chữ “mesh offline”, nhịp tim và thông tin cứu hộ trong mock UI không phải tính năng đã được tích hợp.')
h('Bàn giao tối thiểu')
p('Mã mobile/backend, contract request/response, migration mới nếu cần, hướng dẫn cấu hình không chứa secret, dữ liệu demo và biên bản kiểm thử luồng GPS → deviation → checkpoint → inbox. Không thay đổi ownership của Trip/Trail/Checkpoint definition.')

page('6 Kiểm thử và các quyết định cần chốt')
h('Kịch bản nghiệm thu xuyên suốt')
p('1. User tham gia trip hợp lệ đăng nhập, mở map thật, cấp quyền GPS và bắt đầu session. Route/checkpoint phải đúng trip; người ngoài trip không xem được vị trí.')
p('2. Đi một đoạn thực tế: marker thay đổi, actual path tăng, batch lưu thành công. Gửi lại batch không tăng số mẫu; mẫu accuracy kém hoặc tọa độ sai không được dùng để tự đánh dấu checkpoint.')
p('3. Tắt mạng rồi đi tiếp, bật mạng lại: local path còn, queue gửi bù không trùng. Khóa màn hình theo phạm vi đã chốt và kiểm tra rõ hệ thống còn thu hay chuyển sang stale.')
p('4. Mô phỏng lệch tuyến đủ ngưỡng: user có warning, leader đúng trip có alert/inbox; retry không spam. Quay lại tuyến đóng trạng thái lệch. Đổi active trail phải dùng geometry mới.')
p('5. Vào checkpoint: tạo một arrival, mở mission đúng role, completion cập nhật Mảng A. Dừng/kết thúc chuyến đi đóng session và giữ summary. Thử weather lỗi và consumer restart để kiểm tra phục hồi.')
h('Các quyết định team cần ghi thành contract')
table(['Nội dung','Đề xuất để chốt'],[
 ('Nền tảng MVP','Chọn thiết bị Android/iOS mục tiêu; xác nhận yêu cầu terrain 3D trên native. Không nghiệm thu bằng map mock hoặc chỉ map web.'),
 ('GPS chạy nền','Chốt demo foreground hay tracking khi khóa màn hình. Nếu cần tracking thực địa liên tục, background phải nằm trong P0.'),
 ('Policy GPS','Chốt accuracy tối đa, số mẫu và ngưỡng deviation, interval, radius, stale timeout, retention sau thử thiết bị.'),
 ('Checkpoint','Chốt tripCheckpointId trong API, owner mission-complete, payload event và quy tắc User/Leader/Private Host.'),
 ('Notification','Chỉ định owner dùng chung; xác nhận PostgreSQL hay kế hoạch MongoDB riêng; chốt inbox và yêu cầu push khi app đóng.'),
 ('Độ tin cậy','Chốt thời hạn dữ liệu cũ, backoff/retry và SLA demo leader nhận alert; không gắn SLA này với khả năng nhận push chưa được thử.')
],[1.45,5.25])
p('Ưu tiên bước tiếp theo: làm C1 và C2 thành một bản demo dọc với một user trên một trip, sau đó bổ sung deviation, checkpoint và inbox. Điều kiện đánh dấu hoàn tất là dữ liệu thật đi qua toàn bộ luồng, không chỉ có schema hoặc màn hình.')

page('7 Nguồn đối chiếu')
p('Đường dẫn dưới đây tính từ thư mục gốc TREK_GO. Các nhóm nguồn cho phép truy lại kết luận và rà soát khi code thay đổi.')
sources=[
 ('S1','api/docs/luat.md; api/docs/master.md — ownership, FR, priority, contract và nghiệm thu.'),
 ('S2','web/package.json; web/src/features/core/components/map/BaseMap3D.tsx; map-config.ts; map-utils.ts trong cùng thư mục — thư viện, nguồn tiles, terrain, polyline, marker.'),
 ('S3','web/src/features/leaders/components/leader/TripRouteReviewMap.tsx; StepTwoOverviewMap.tsx trong cùng thư mục — lấy dữ liệu system trail và hiển thị bản đồ.'),
 ('S4','mobile/package.json; mobile/src/features/trip/components/TripMap/index.tsx; mobile/src/features/trip/mockData.ts — dependencies và UI mô phỏng.'),
 ('S5','api/prisma/schema.prisma — GpsPolicy, GpsSession, GpsLocation, RouteDeviationEvent, CheckpointEvent, WeatherSnapshot, GpsOutboxEvent, Notification.'),
 ('S6','api/src/modules/trips/presentation/controllers/trip-navigation.controller.ts; application/use-cases/get-navigation-context.use-case.ts; application/use-cases/switch-trail.use-case.ts trong module trips.'),
 ('S7','api/src/modules/geo/infrastructure/adapters/prisma-trip-map-state.adapter.ts; application/use-cases/get-trip-map-state.use-case.ts; presentation/controllers/trip-map-state.controller.ts trong module geo.'),
 ('S8','api/docs/OUTBOX_PATTERN_EXPLAINED.md — hướng dẫn có ví dụ MongoDB và consumer; cần đồng bộ lại với cây nguồn.'),
 ('S9','api/src/app.module.ts; api/package.json; api/src/shared/config.ts; web/src/features/leaders/components/leader/LeaderSafetySettings.tsx; web/src/features/rentals/components/staff/StaffProfile.tsx.'),
 ('S10','api/src/modules/bookings/booking-cancellation.service.ts — dispatch hoàn tiền, interval và đánh dấu outbox PROCESSED.'),
 ('S11','api/src/modules/identity/presentation/controllers/auth.controller.ts; api/src/modules/identity/infrastructure/adapters/supabase-identity.adapter.ts — auth hiện hành cần tái sử dụng.'),
]
for key,val in sources:p(key+'  '+val)
h('Tài liệu thư viện chính thức')
p('W1  Expo Location SDK 57 — quyền và giới hạn foreground của watchPositionAsync\nhttps://docs.expo.dev/versions/v57.0.0/sdk/location/')
p('W2  MapLibre React Native — setup native và rebuild\nhttps://maplibre.org/maplibre-react-native/docs/setup/getting-started/')
p('W3  MapLibre GL JS GeolocateControl — nguồn geolocation cho web\nhttps://maplibre.org/maplibre-gl-js/docs/API/classes/GeolocateControl/')
d.core_properties.title='GPS bản đồ và thông báo trong MVP của Luật'
d.core_properties.subject='TrekGo Area C implementation review and MVP backlog'
d.save(OUT)
print(OUT)
